import re

from pathlib import Path
import os, sys, re
import tempfile
from datetime import datetime, timedelta
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

try:
    from PIL import Image, ImageDraw, ImageFont, ImageChops
except Exception:
    Image = ImageDraw = ImageFont = None

from platform_utils import normalize_input_file, setup_logger

try:
    from openpyxl import load_workbook
except ImportError:
    raise SystemExit("Please install dependencies: python3 -m pip install python-docx openpyxl")

def convert_numbers_to_xlsx(numbers_path):
    """
    Convert Apple Numbers file to XLSX on macOS using the Numbers app.
    This requires Apple Numbers installed.
    """
    import subprocess
    import tempfile
    import platform

    numbers_path = Path(numbers_path).resolve()
    if platform.system().lower() != "darwin":
        raise RuntimeError(".numbers input is supported only on macOS with Apple Numbers installed.")

    if not numbers_path.exists():
        raise RuntimeError(f"Numbers file not found: {numbers_path}")

    temp_dir = Path(tempfile.mkdtemp(prefix="proposal_numbers_"))
    xlsx_out = temp_dir / (numbers_path.stem + ".xlsx")

    applescript = f"""
set inputFile to POSIX file "{numbers_path}" as alias
set outputFile to POSIX file "{xlsx_out}"
tell application "Numbers"
    activate
    open inputFile
    delay 1
    set theDoc to front document
    export theDoc to outputFile as Microsoft Excel
    close theDoc saving no
end tell
"""

    result = subprocess.run(["osascript", "-e", applescript], capture_output=True, text=True)
    if result.returncode != 0:
        raise RuntimeError(
            "Unable to convert .numbers file to .xlsx. Please ensure Apple Numbers is installed and permissions are allowed.\n"
            + (result.stderr or result.stdout)
        )

    if not xlsx_out.exists():
        raise RuntimeError("Numbers conversion completed but XLSX file was not created.")

    return xlsx_out

def normalize_input_excel(input_path):
    input_path = Path(input_path)
    if input_path.suffix.lower() == ".numbers":
        return convert_numbers_to_xlsx(input_path)
    if input_path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise RuntimeError("Input must be .xlsx, .xlsm, or .numbers")
    return input_path

NAVY = "0D2742"
BLUE = "0F5FAD"
LIGHT_BLUE = "EAF2FB"
WHITE = "FFFFFF"
TEXT = "1F2937"

def clean(v):
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() in ("none", "nan") else s

def parse_number_or_zero(v):
    """Return 0 for missing, blank, deleted, invalid or zero GST values."""
    if v is None:
        return 0.0
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return 0.0
        s = s.replace("%", "").replace(",", "").strip()
        if not s:
            return 0.0
        try:
            n = float(s)
            return n / 100.0 if n > 1 else n
        except Exception:
            return 0.0
    try:
        n = float(v)
        return n / 100.0 if n > 1 else n
    except Exception:
        return 0.0

def normalize_header(v):
    """Normalize Excel headers like 'GST %', 'GST%', 'GST (%)', 'GST Rate'."""
    return re.sub(r"[^a-z0-9]+", "", clean(v).lower())

def is_gst_header(v):
    h = normalize_header(v)
    return h in ("gst", "gstpercent", "gstpct", "gstrate", "gstpercentage") or ("gst" in h)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_border(cell, color="D8E2EE", sz="8"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        el = tcBorders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            tcBorders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)

def set_cell_text(cell, text, bold=False, color=TEXT, size=10.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(clean(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_twips):
    """Set a fixed table width in twentieths of a point."""
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(width_twips))
    tblW.set(qn("w:type"), "dxa")


def set_table_column_widths(table, widths_inches):
    """Set fixed Word table layout and practical column widths."""
    table.autofit = False
    tblPr = table._tbl.tblPr

    # Force fixed layout so Word/LibreOffice does not redistribute columns equally.
    tblLayout = tblPr.find(qn("w:tblLayout"))
    if tblLayout is None:
        tblLayout = OxmlElement("w:tblLayout")
        tblPr.append(tblLayout)
    tblLayout.set(qn("w:type"), "fixed")

    total_twips = int(sum(widths_inches) * 1440)
    set_table_width(table, total_twips)

    # Define table grid
    tblGrid = table._tbl.tblGrid
    if tblGrid is not None:
        table._tbl.remove(tblGrid)
    tblGrid = OxmlElement("w:tblGrid")
    table._tbl.insert(0, tblGrid)
    for width in widths_inches:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 1440)))
        tblGrid.append(grid_col)

    for row in table.rows:
        for idx, width in enumerate(widths_inches):
            if idx >= len(row.cells):
                continue
            cell = row.cells[idx]
            cell.width = Inches(width)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(width * 1440)))
            tcW.set(qn("w:type"), "dxa")

def style_table(table, header_fill=NAVY, zebra=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, header_fill)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.color.rgb = RGBColor.from_string(WHITE)
                        run.bold = True
                        run.font.name = "Aptos"
            elif zebra and r_idx % 2 == 0:
                set_cell_shading(cell, "F3F6FA")
            else:
                set_cell_shading(cell, WHITE)

def resolve_asset_path(xlsx_path, configured_path, fallback_name):
    """Resolve asset paths from template folder or quotation module folder.

    Earlier versions resolved Logo Path relative only to the templates folder,
    which caused the logo to be missed when assets were stored in ../assets.
    """
    xlsx_path = Path(xlsx_path)
    configured = clean(configured_path) or f"assets/{fallback_name}"
    candidates = []
    p = Path(configured)
    if p.is_absolute():
        candidates.append(p)
    else:
        candidates.append(xlsx_path.parent / configured)
        candidates.append(Path(__file__).parent / configured)
        candidates.append(Path(__file__).parent / "assets" / Path(configured).name)
        candidates.append(xlsx_path.parent.parent / configured)
        candidates.append(xlsx_path.parent.parent / "assets" / Path(configured).name)
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return candidates[0]


def add_page_header(section, logo_path, company, tagline):
    header = section.header
    header.is_linked_to_previous = False
    for p in list(header.paragraphs):
        p._element.getparent().remove(p._element)

    table = header.add_table(rows=1, cols=2, width=Inches(7.1))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Inches(1.25)
    table.columns[1].width = Inches(5.85)

    for cell in table.row_cells(0):
        set_cell_border(cell, "FFFFFF", "0")
        set_cell_margins(cell, 0, 0, 0, 0)

    left, right = table.row_cells(0)

    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if Path(logo_path).exists():
        run = p.add_run()
        run.add_picture(str(logo_path), width=Inches(0.92))

    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(clean(company))
    r.bold = True
    r.font.size = Pt(12.5)
    r.font.name = "Aptos Display"
    r.font.color.rgb = RGBColor.from_string(BLUE)

    p2 = right.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p2.add_run(clean(tagline))
    r.font.size = Pt(9.7)
    r.font.color.rgb = RGBColor.from_string("6B7280")
    r.font.name = "Aptos"

    # Subtle branding separator in the header
    p3 = right.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p3.add_run("Commercial Document Suite")
    r.font.size = Pt(8.2)
    r.font.color.rgb = RGBColor.from_string("8A97A8")
    r.font.name = "Aptos"

def _pick_setting(settings, *keys):
    for key in keys:
        value = clean(settings.get(key))
        if value:
            return value
    return ""


def _pick_customer(customer, *keys):
    for key in keys:
        value = clean(customer.get(key))
        if value:
            return value
    return ""


def _add_plain_line(cell, text, size=8.2, bold=False, color=TEXT, after=1):
    if not clean(text):
        return
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(clean(text))
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)


def _add_label_value_line(cell, label, value, size=8.2):
    if not clean(value):
        return
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(label + ": ")
    r.bold = True
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(TEXT)
    r = p.add_run(clean(value))
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(TEXT)



def _font(size, bold=False):
    """Return a scalable invoice-style font.

    Important fix:
    On macOS the previous font lookup often missed Arial/Helvetica, so Pillow
    fell back to a tiny bitmap default font. That is why the quotation header
    text became unreadably small on GUI-generated reports.
    """
    if ImageFont is None:
        return None

    candidates = [
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",

        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
        "/System/Library/Fonts/Supplemental/Helvetica.ttf",

        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
    ]

    for fp in candidates:
        try:
            if Path(fp).exists():
                return ImageFont.truetype(fp, size)
        except Exception:
            pass

    search_roots = [
        Path("/System/Library/Fonts"),
        Path("/System/Library/Fonts/Supplemental"),
        Path("/Library/Fonts"),
        Path.home() / "Library" / "Fonts",
        Path("C:/Windows/Fonts"),
        Path("/usr/share/fonts"),
    ]
    preferred = [
        "Arial Bold.ttf" if bold else "Arial.ttf",
        "LiberationSans-Bold.ttf" if bold else "LiberationSans-Regular.ttf",
        "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf",
        "Helvetica.ttc",
        "Calibri Bold.ttf" if bold else "Calibri.ttf",
        "Segoe UI Bold.ttf" if bold else "Segoe UI.ttf",
    ]
    for root in search_roots:
        try:
            if not root.exists():
                continue
            for name in preferred:
                matches = list(root.rglob(name))
                if matches:
                    return ImageFont.truetype(str(matches[0]), size)
        except Exception:
            continue

    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()

def _draw_text_fit(draw, xy, text, font_size, max_width, fill, bold=True, min_size=17):
    """Draw text, reducing font size until it fits the available width."""
    text = clean(text)
    size = font_size
    while size >= min_size:
        f = _font(size, bold)
        bbox = draw.textbbox((0, 0), text, font=f)
        if bbox[2] - bbox[0] <= max_width:
            draw.text(xy, text, font=f, fill=fill)
            return size
        size -= 1
    f = _font(min_size, bold)
    draw.text(xy, text, font=f, fill=fill)
    return min_size


def _crop_logo_padding(logo):
    """Crop transparent/near-white padding around logo without damaging visible ring."""
    logo = logo.convert("RGBA")
    alpha_bbox = logo.getbbox()
    if alpha_bbox:
        logo = logo.crop(alpha_bbox)
    # crop near-white background padding
    rgb_white = Image.new("RGBA", logo.size, (255, 255, 255, 255))
    diff = ImageChops.difference(logo, rgb_white)
    bbox = diff.getbbox()
    if bbox:
        logo = logo.crop(bbox)
    return logo


def _render_invoice_style_quotation_header_image(logo_path, settings, quotation, customer, out_path):
    """Invoice-code-matched quotation header.

    This uses the same visual geometry as the invoice header:
    - navy top strip
    - left company name, right document title
    - logo on left
    - company information in center
    - compact details card on right
    """
    if Image is None:
        return None

    company = _pick_setting(settings, "Company Name") or "SMS Controls & Automation"
    tagline = _pick_setting(settings, "Tagline") or "Industrial Automation | AC/DC Drives | Repairs | Retrofitting"
    address1 = _pick_setting(settings, "Address Line 1", "Address", "Company Address") or "Office No.2, Plot No.241, Loha Mandi"
    address2 = _pick_setting(settings, "Address Line 2") or "Ghaziabad, Uttar Pradesh - 201009"
    email = _pick_setting(settings, "Email") or "info@smscontrols.com"
    phone = _pick_setting(settings, "Phone", "Mobile", "Contact") or "+91-8826059159 / 8595231536"
    gstin = _pick_setting(settings, "GSTIN", "GST No.", "GST") or "09JSMPS2386Q1ZB"
    state = _pick_setting(settings, "State") or "Uttar Pradesh"
    state_code = _pick_setting(settings, "State Code", "Code") or "09"

    def _display(v):
        if hasattr(v, "strftime"):
            return v.strftime("%d-%b-%Y")
        return clean(v)

    proposal_no = _display(quotation.get("Quotation No.")) or _display(quotation.get("Quotation No")) or _display(quotation.get("Proposal No")) or "-"
    proposal_date = _display(quotation.get("Date")) or "-"
    validity = _display(quotation.get("Validity")) or "-"

    # Logical canvas uses same proportions as the working invoice header screenshot.
    SCALE = 3
    BW, BH = 1500, 350
    W, H = BW * SCALE, BH * SCALE
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)

    navy = "#0D344F"
    blue = "#1E96D4"
    soft = "#EAF2FB"
    text_col = "#0B2545"
    muted = "#5B6B82"
    white = "#FFFFFF"

    def sc(v):
        return int(round(v * SCALE))

    def font(size, bold=False):
        return _font(sc(size), bold)

    def rect(t):
        return tuple(sc(v) for v in t)

    # Outer container and top strip.
    d.rounded_rectangle(rect((2, 2, BW-3, BH-3)), radius=sc(7), outline=navy, width=sc(2), fill=white)
    d.rectangle(rect((2, 2, BW-3, 72)), fill=navy)

    d.text((sc(32), sc(25)), company, font=font(31, True), fill=white)
    doc_title = "TECHNO-COMMERCIAL PROPOSAL"
    tf = font(29, True)
    tb = d.textbbox((0, 0), doc_title, font=tf)
    d.text((W - sc(32) - (tb[2] - tb[0]), sc(27)), doc_title, font=tf, fill=white)

    # Logo - invoice-like position and size.
    logo_box = rect((68, 108, 288, 308))
    lp = Path(logo_path)
    if lp.exists():
        try:
            logo = _crop_logo_padding(Image.open(lp))
            target_w = logo_box[2] - logo_box[0]
            target_h = logo_box[3] - logo_box[1]
            scale = min(target_w / logo.width, target_h / logo.height)
            logo = logo.resize((max(1, int(logo.width * scale)), max(1, int(logo.height * scale))), Image.LANCZOS)
            lx = logo_box[0] + (target_w - logo.width) // 2
            ly = logo_box[1] + (target_h - logo.height) // 2
            img.paste(logo, (lx, ly), logo)
        except Exception:
            d.text((sc(110), sc(190)), "SMS", font=font(62, True), fill=navy)
    else:
        d.text((sc(110), sc(190)), "SMS", font=font(62, True), fill=navy)

    # Company block - same balance as invoice.
    cx = sc(335)
    _draw_text_fit(d, (cx, sc(102)), company, sc(42), sc(610), navy, True, sc(32))
    d.text((cx, sc(154)), tagline, font=font(23, False), fill=muted)
    d.text((cx, sc(198)), clean(address1), font=font(21, False), fill=text_col)
    d.text((cx, sc(233)), clean(address2), font=font(21, False), fill=text_col)
    d.text((cx, sc(268)), f"GSTIN: {gstin} | State: {state}, Code: {state_code}", font=font(20, False), fill=text_col)
    d.text((cx, sc(303)), f"Email: {email} | Phone: {phone}", font=font(20, False), fill=text_col)

    # Proposal details card.
    card_w, card_h = sc(440), sc(205)
    card_x, card_y = W - sc(72) - card_w, sc(95)
    d.rounded_rectangle((card_x, card_y, card_x + card_w, card_y + card_h), radius=sc(6), outline=blue, width=sc(2), fill=soft)
    d.rectangle((card_x, card_y, card_x + card_w, card_y + sc(58)), fill=navy)
    head = "PROPOSAL DETAILS"
    hf = font(27, True)
    hb = d.textbbox((0, 0), head, font=hf)
    d.text((card_x + (card_w - (hb[2]-hb[0]))//2, card_y + sc(17)), head, font=hf, fill=white)

    rows = [("Proposal No.", proposal_no), ("Proposal Date", proposal_date), ("Validity", validity)]
    y = card_y + sc(88)
    for label, value in rows:
        d.text((card_x + sc(38), y), label, font=font(20, True), fill=muted)
        d.text((card_x + sc(205), y), ":", font=font(20, True), fill=text_col)
        d.text((card_x + sc(235), y), value, font=font(20, True if label == "Proposal No." else False), fill=text_col)
        y += sc(43)

    img.save(out_path)
    return out_path

def add_first_page_branding(doc, logo_path, settings, quotation, customer):
    """Add quotation branding header using the same invoice-style rendered geometry."""
    temp_header = Path(tempfile.gettempdir()) / f"sms_quotation_invoice_style_header_{os.getpid()}_{int(datetime.now().timestamp())}.png"
    rendered = _render_invoice_style_quotation_header_image(logo_path, settings, quotation, customer, temp_header)
    if rendered and Path(rendered).exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(8)
        p.add_run().add_picture(str(rendered), width=Inches(7.25))
        try:
            Path(rendered).unlink(missing_ok=True)
        except Exception:
            pass
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SMS Controls & Automation")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TECHNO-COMMERCIAL PROPOSAL")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    doc.add_paragraph("")

def add_proposal_cover_section(doc, settings, quotation, customer):
    """Restore proposal title, product subtitle, and Prepared For/Prepared By block."""
    title = clean(quotation.get("Proposal Title")) or "TECHNO-COMMERCIAL PROPOSAL"
    product_title = clean(quotation.get("Product Title"))
    product_subtitle = clean(quotation.get("Product Subtitle"))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor.from_string(NAVY)

    if product_title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(product_title)
        r.bold = True
        r.font.name = "Aptos Display"
        r.font.size = Pt(19)
        r.font.color.rgb = RGBColor.from_string(BLUE)

    if product_subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(product_subtitle)
        r.bold = True
        r.font.name = "Aptos Display"
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor.from_string("6B7280")

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    # V2.8: use the full available page width, matching the invoice/header width.
    set_table_width(table, 10580)
    table.columns[0].width = Inches(3.62)
    table.columns[1].width = Inches(3.62)
    for cell in table.rows[0].cells:
        cell.width = Inches(3.62)
    left, right = table.rows[0].cells
    for i, cell in enumerate([left, right]):
        set_cell_border(cell, "BBD2EA", "8")
        set_cell_shading(cell, "EAF2FB" if i == 0 else WHITE)
        set_cell_margins(cell, 210, 180, 210, 180)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    left.text = ""
    p = left.paragraphs[0]
    r = p.add_run("PREPARED FOR")
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    cust_name = clean(customer.get("Customer Name"))
    to_name = clean(customer.get("To"))
    address = clean(customer.get("Address"))
    gstin = clean(customer.get("GSTIN"))
    email = clean(customer.get("Email"))
    mobile = clean(customer.get("Mobile"))

    if cust_name:
        p = left.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        r = p.add_run(cust_name)
        r.bold = True
        r.font.name = "Aptos Display"
        r.font.size = Pt(17)
        r.font.color.rgb = RGBColor.from_string(TEXT)
    for label, value in [("To", to_name), ("Address", address), ("GSTIN", gstin), ("Mobile", mobile), ("Email", email)]:
        if value:
            p = left.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(label + ": ")
            r.bold = True
            r.font.name = "Aptos"
            r.font.size = Pt(12.2)
            r.font.color.rgb = RGBColor.from_string(TEXT)
            r = p.add_run(value)
            r.font.name = "Aptos"
            r.font.size = Pt(12.2)
            r.font.color.rgb = RGBColor.from_string(TEXT)

    right.text = ""
    p = right.paragraphs[0]
    r = p.add_run("PREPARED BY")
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    company = _pick_setting(settings, "Company Name") or clean(quotation.get("Prepared By")) or "SMS Controls & Automation"
    tagline = "Industrial Automation | Drives | Panels | PLC | SCADA | Services | Retrofits"
    p = right.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run(company)
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(17)
    r.font.color.rgb = RGBColor.from_string(TEXT)
    p = right.add_paragraph()
    r = p.add_run(str(tagline).replace("\n", " "))
    r.font.name = "Aptos"
    r.font.size = Pt(9.4)
    r.font.color.rgb = RGBColor.from_string(TEXT)

    # Keep a balanced fixed height similar to the original prepared-for section.
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trHeight = OxmlElement("w:trHeight")
        trHeight.set(qn("w:val"), "2100")
        trHeight.set(qn("w:hRule"), "atLeast")
        trPr.append(trHeight)
    doc.add_paragraph("")

def add_footer(section, footer_text):
    footer = section.footer
    footer.is_linked_to_previous = False
    for p in list(footer.paragraphs):
        p._element.getparent().remove(p._element)
    p = footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(clean(footer_text))
    r.font.size = Pt(8.5)
    r.font.name = "Aptos"
    r.font.color.rgb = RGBColor.from_string("6B7280")

def heading(doc, title, subtitle=None):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.columns[0].width = Inches(0.14)
    table.columns[1].width = Inches(6.9)
    c0, c1 = table.row_cells(0)
    set_cell_shading(c0, BLUE)
    set_cell_border(c0, BLUE, "0")
    set_cell_border(c1, WHITE, "0")
    p = c1.paragraphs[0]
    r = p.add_run(clean(title))
    r.bold = True
    r.font.size = Pt(18)
    r.font.name = "Aptos Display"
    r.font.color.rgb = RGBColor.from_string(NAVY)
    if subtitle:
        p2 = c1.add_paragraph()
        r = p2.add_run(clean(subtitle))
        r.font.size = Pt(11.5)
        r.font.color.rgb = RGBColor.from_string("6B7280")
        r.font.name = "Aptos"
    doc.add_paragraph("")

def fmt_inr(n):
    try:
        n = int(round(float(n)))
    except Exception:
        n = 0
    s = str(abs(n))
    if len(s) <= 3:
        out = s
    else:
        out = s[-3:]
        s = s[:-3]
        while len(s) > 2:
            out = s[-2:] + "," + out
            s = s[:-2]
        if s:
            out = s + "," + out
    return ("-" if n < 0 else "") + "₹" + out

def fmt_date(value):
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, (int, float)):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).strftime("%d.%m.%Y")
        except Exception:
            return clean(value)
    return clean(value)

ONES = ["","One","Two","Three","Four","Five","Six","Seven","Eight","Nine","Ten","Eleven","Twelve","Thirteen","Fourteen","Fifteen","Sixteen","Seventeen","Eighteen","Nineteen"]
TENS = ["","","Twenty","Thirty","Forty","Fifty","Sixty","Seventy","Eighty","Ninety"]

def two_digit(n):
    n = int(n)
    if n < 20:
        return ONES[n]
    return TENS[n//10] + ("" if n % 10 == 0 else " " + ONES[n%10])

def three_digit(n):
    n = int(n)
    if n < 100:
        return two_digit(n)
    return ONES[n//100] + " Hundred" + ("" if n % 100 == 0 else " " + two_digit(n%100))

def amount_words(n):
    n = int(round(float(n)))
    if n == 0:
        return "Zero Rupees Only"
    parts = []
    crore = n // 10000000
    n %= 10000000
    lakh = n // 100000
    n %= 100000
    thousand = n // 1000
    n %= 1000
    if crore:
        parts.append(three_digit(crore) + " Crore")
    if lakh:
        parts.append(three_digit(lakh) + " Lakh")
    if thousand:
        parts.append(three_digit(thousand) + " Thousand")
    if n:
        parts.append(three_digit(n))
    return "Rupees " + " ".join(parts) + " Only"

def get_map(ws):
    data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None:
            data[clean(row[0])] = row[1]
    return data

def read_excel(xlsx_path):
    # wb_formula keeps headers/text; wb_values reads displayed/calculated Excel values.
    # This is important when GST % cells show 18% but are actually formulas.
    wb_formula = load_workbook(xlsx_path, data_only=False)
    wb_values = load_workbook(xlsx_path, data_only=True)
    customer = get_map(wb_values["Customer_Details"])
    quotation = get_map(wb_values["Quotation_Details"])
    static = get_map(wb_values["Static_Content"])
    settings = get_map(wb_values["Settings"])
    terms = get_map(wb_values["Terms"])

    items = []
    items_ws = wb_formula["Items"]
    items_values_ws = wb_values["Items"]

    # Dynamic column handling with robust GST detection.
    # Supports GST headers like: GST %, GST%, GST (%), GST Rate, GST Percent, Tax GST, etc.
    header_values = [clean(c.value) for c in items_ws[1]]
    normalized_headers = [normalize_header(h) for h in header_values]

    def find_col(*names):
        wanted = {normalize_header(n) for n in names}
        for idx, h in enumerate(normalized_headers):
            if h in wanted:
                return idx
        return None

    def find_gst_col():
        for idx, raw in enumerate(header_values):
            if is_gst_header(raw):
                return idx
        return None

    idx_sr = find_col("Sr. No.", "Sr No", "S.No", "Sr")
    idx_desc = find_col("Description", "Item Description", "Product", "Product Description")
    idx_qty = find_col("Qty", "Quantity")
    idx_unit = find_col("Unit", "UOM")
    idx_unit_price = find_col("Unit Price", "Price", "Rate")
    idx_disc = find_col("Discount", "Discount %", "Discount%", "Disc %", "Disc")
    idx_gst = find_gst_col()

    # Safe fallback for old V10 layout:
    # Sr, Description, Qty, Unit, Unit Price, Discount, GST %
    if idx_sr is None:
        idx_sr = 0
    if idx_desc is None:
        idx_desc = 1
    if idx_qty is None:
        idx_qty = 2
    if idx_unit is None:
        idx_unit = 3
    if idx_unit_price is None:
        idx_unit_price = 4

    gst_column_exists = idx_gst is not None
    item_gst_amounts = []
    item_gst_values = []

    def get_cell(row, idx):
        return row[idx] if idx is not None and idx < len(row) else None

    for formula_row, value_row in zip(
        items_ws.iter_rows(min_row=2, max_row=50, values_only=False),
        items_values_ws.iter_rows(min_row=2, max_row=50, values_only=False)
    ):
        row = [vc.value if vc.value is not None else fc.value for fc, vc in zip(formula_row, value_row)]
        sr = get_cell(row, idx_sr)
        desc = get_cell(row, idx_desc)
        qty = get_cell(row, idx_qty)
        unit = get_cell(row, idx_unit)
        unit_price = get_cell(row, idx_unit_price)
        disc = get_cell(row, idx_disc)
        gst_value = get_cell(row, idx_gst)

        if not clean(desc):
            continue

        qty = float(qty or 0)
        unit_price = float(unit_price or 0)
        disc = parse_number_or_zero(disc)

        if qty <= 0:
            raise ValueError(f"Quantity must be greater than zero for item: {desc}")
        if disc < 0:
            raise ValueError(f"Discount cannot be negative for item: {desc}")

        gross = qty * unit_price
        discount_amount = gross * disc
        net = gross - discount_amount

        gst_rate = parse_number_or_zero(gst_value) if gst_column_exists else 0.0
        if gst_column_exists:
            item_gst_values.append(gst_rate)
            item_gst_amounts.append(net * gst_rate)

        items.append({
            "sr": int(sr) if sr else len(items)+1,
            "description": clean(desc),
            "qty": qty,
            "unit": clean(unit),
            "unit_price": unit_price,
            "discount_pct": disc,
            "gross": gross,
            "discount_amount": discount_amount,
            "net": net,
            "gst_pct": gst_rate
        })
    if not items:
        raise ValueError("No quotation items found in Items sheet.")

    original = sum(i["gross"] for i in items)
    discount = sum(i["discount_amount"] for i in items)
    final = sum(i["net"] for i in items)

    # GST decision rule:
    # 1. If Items sheet has a GST column, GST is decided from that column only.
    #    If all GST cells are blank/deleted/0, GST is NOT applied even if Settings has Default GST %.
    # 2. If Items sheet has no GST column, fallback to Settings -> Default GST %.
    if gst_column_exists:
        gst_amount = sum(item_gst_amounts)
        gst_pct = (gst_amount / final) if final else 0.0
    else:
        gst_pct = parse_number_or_zero(settings.get("Default GST %"))
        if gst_pct < 0:
            raise ValueError("GST percentage cannot be negative.")
        gst_amount = final * gst_pct

    grand_total = final + gst_amount
    has_discount = any(abs(i["discount_pct"]) > 0.000001 for i in items)
    return customer, quotation, static, settings, terms, items, {
        "original": original,
        "discount": discount,
        "final": final,
        "gst_pct": gst_pct,
        "gst_amount": gst_amount,
        "grand_total": grand_total,
        "words": amount_words(final),
        "has_discount": has_discount
    }

def add_if_present(cell, label, value, bold_label=True):
    value = clean(value)
    if not value:
        return
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1
    if label:
        r = p.add_run(label + ": ")
        r.bold = bold_label
        r.font.name = "Aptos"
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor.from_string(TEXT)
    r = p.add_run(value)
    r.font.name = "Aptos"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(TEXT)

def add_intro_cards(doc, static):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cards = [
        (static.get("Experience Metric","15+"), static.get("Experience Label","Years Experience")),
        (static.get("Drive Metric","AC/DC"), static.get("Drive Label","Drive Specialist")),
        (static.get("Tech Metric","PLC/HMI"), static.get("Tech Label","SCADA Expertise")),
        (static.get("Support Metric","Fast"), static.get("Support Label","Technical Support")),
    ]
    for i, (metric, label) in enumerate(cards):
        c = table.cell(0, i)
        set_cell_shading(c, LIGHT_BLUE)
        set_cell_border(c, "D8E2EE")
        set_cell_margins(c, 120, 80, 120, 80)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(clean(metric))
        r.bold=True
        r.font.size=Pt(17)
        r.font.color.rgb=RGBColor.from_string(BLUE)
        r.font.name="Aptos Display"
        p2 = c.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p2.add_run(clean(label))
        r.bold=True
        r.font.size=Pt(9.5)
        r.font.color.rgb=RGBColor.from_string(TEXT)
        r.font.name="Aptos"

def create_docx(xlsx_path, output_path):
    xlsx_path = Path(xlsx_path)
    base = xlsx_path.parent
    customer, quotation, static, settings, terms, items, totals = read_excel(xlsx_path)
    logo_path = resolve_asset_path(xlsx_path, settings.get("Logo Path", "assets/logo_from_doc.png"), "logo_from_doc.png")
    industry_path = resolve_asset_path(xlsx_path, settings.get("Industry Image Path", "assets/industry_collage.png"), "industry_collage.png")

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.45)
    sec.bottom_margin = Inches(0.48)
    sec.left_margin = Inches(0.55)
    sec.right_margin = Inches(0.55)
    # Do not use a Word page header for quotation. Branding is placed only
    # on first page body; continuation pages keep footer only.
    add_footer(sec, static.get("Footer", f'{settings.get("Company Name")} | {settings.get("Email")} | {settings.get("Mobile")}'))

    # Page 1 premium branding block: same visual style as invoice header.
    add_first_page_branding(doc, logo_path, settings, quotation, customer)

    # Restore original proposal cover content before the salutation.
    add_proposal_cover_section(doc, settings, quotation, customer)

    intro = doc.add_table(rows=1, cols=1)
    intro.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = intro.cell(0,0)
    set_cell_shading(c, "F7F9FC")
    set_cell_border(c)
    set_cell_margins(c, 120,120,120,120)
    if clean(static.get("Salutation")):
        set_cell_text(c, static.get("Salutation"), True, TEXT, 11)
    if clean(static.get("Intro Text")):
        p = c.add_paragraph(clean(static.get("Intro Text")))
        p.runs[0].font.name="Aptos"
        p.runs[0].font.size=Pt(10.5)
    doc.add_paragraph("")
    add_intro_cards(doc, static)

    doc.add_page_break()

    # Page 2
    heading(doc, "COMPANY PROFILE & OFFERED SOLUTION", "Compact overview of capabilities and proposed supply")
    about = doc.add_table(rows=1, cols=2)
    about.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, cell in enumerate(about.row_cells(0)):
        set_cell_border(cell)
        set_cell_shading(cell, LIGHT_BLUE if idx==1 else WHITE)
        set_cell_margins(cell, 120,120,120,120)
    left, right = about.row_cells(0)
    set_cell_text(left, "ABOUT SMS CONTROLS & AUTOMATION", True, BLUE, 10.5)
    for txt in [static.get("Company Overview"), static.get("Industry Text")]:
        if clean(txt):
            p = left.add_paragraph(clean(txt))
            p.runs[0].font.size=Pt(10.5)
            p.runs[0].font.name="Aptos"
    set_cell_text(right, "CORE CAPABILITIES", True, BLUE, 10.5)
    for i in range(1, 7):
        cap = clean(static.get(f"Capability {i}"))
        if cap:
            p = right.add_paragraph("✓ " + cap)
            p.runs[0].font.size=Pt(10.5)
            p.runs[0].font.name="Aptos"
    if industry_path.exists():
        p = right.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(industry_path), width=Inches(2.85))

    doc.add_paragraph("")
    heading(doc, "SCOPE OF SUPPLY")
    table = doc.add_table(rows=1, cols=3)
    set_table_column_widths(table, [0.72, 5.68, 0.85])
    for i, h in enumerate(["Sr. No.", "Description", "Qty"]):
        set_cell_text(table.cell(0,i), h, True, WHITE, 10.5, WD_ALIGN_PARAGRAPH.CENTER)
    for it in items:
        row = table.add_row().cells
        set_cell_text(row[0], f"{it['sr']:02d}", False, TEXT, 10)
        set_cell_text(row[1], it["description"], False, TEXT, 10)
        qty_display = f"{int(it['qty']) if it['qty'].is_integer() else it['qty']} {it['unit']}".strip()
        set_cell_text(row[2], qty_display, False, TEXT, 10, WD_ALIGN_PARAGRAPH.RIGHT)
    set_table_column_widths(table, [0.72, 5.68, 0.85])
    style_table(table)

    doc.add_page_break()

    # Page 3: Dynamic discount table
    heading(doc, "COMMERCIAL OFFER", "Investment summary with automatic calculations from Excel")
    if totals["has_discount"]:
        heads = ["Description", "Qty", "Unit Price", "Discount", "Final Price"]
    else:
        heads = ["Description", "Qty", "Unit Price", "Price"]
    table = doc.add_table(rows=1, cols=len(heads))
    if totals["has_discount"]:
        set_table_column_widths(table, [3.95, 0.70, 1.00, 0.75, 0.85])
    else:
        set_table_column_widths(table, [4.50, 0.75, 1.00, 1.00])
    for i, h in enumerate(heads):
        set_cell_text(table.cell(0,i), h, True, WHITE, 10.3, WD_ALIGN_PARAGRAPH.CENTER)
    for it in items:
        row = table.add_row().cells
        set_cell_text(row[0], it["description"], False, TEXT, 10)
        set_cell_text(row[1], f"{int(it['qty']) if it['qty'].is_integer() else it['qty']} {it['unit']}", False, TEXT, 10)
        set_cell_text(row[2], fmt_inr(it["unit_price"]), False, TEXT, 10, WD_ALIGN_PARAGRAPH.RIGHT)
        if totals["has_discount"]:
            set_cell_text(row[3], f"{it['discount_pct']:.0%}", False, TEXT, 10, WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(row[4], fmt_inr(it["net"]), False, TEXT, 10, WD_ALIGN_PARAGRAPH.RIGHT)
        else:
            set_cell_text(row[3], fmt_inr(it["net"]), False, TEXT, 10, WD_ALIGN_PARAGRAPH.RIGHT)
    if totals["has_discount"]:
        set_table_column_widths(table, [3.95, 0.70, 1.00, 0.75, 0.85])
    else:
        set_table_column_widths(table, [4.50, 0.75, 1.00, 1.00])
    style_table(table)

    doc.add_paragraph("")
    summary = doc.add_table(rows=1, cols=2)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    a, b = summary.row_cells(0)
    set_cell_shading(a, LIGHT_BLUE)
    set_cell_border(a)
    set_cell_margins(a,160,140,160,140)
    set_cell_shading(b, NAVY)
    set_cell_border(b)
    set_cell_margins(b,160,140,160,140)
    set_cell_text(a, "INVESTMENT SUMMARY", True, BLUE, 12)

    has_gst = abs(totals.get("gst_amount", 0)) > 0.000001 and abs(totals.get("gst_pct", 0)) > 0.000001
    if totals["has_discount"]:
        summary_rows = [
            ("Original Offer Value", totals["original"]),
            ("Special Discount", totals["discount"]),
        ]
    else:
        summary_rows = [
            ("Offer Value", totals["final"]),
        ]
    if has_gst:
        summary_rows.extend([
            (f"GST @ {totals['gst_pct']:.0%}", totals["gst_amount"]),
            ("Grand Total Incl. GST", totals["grand_total"]),
        ])
    else:
        summary_rows.append(("Grand Total", totals["final"]))
    for lbl, val in summary_rows:
        p = a.add_paragraph()
        r = p.add_run(lbl + ": ")
        r.bold=True
        r.font.size=Pt(10.5)
        r.font.name="Aptos"
        r.font.color.rgb=RGBColor.from_string(TEXT)
        r = p.add_run(fmt_inr(val))
        r.font.size=Pt(10.5)
        r.font.name="Aptos"

    if has_gst:
        final_heading = "TOTAL PAYABLE VALUE"
        highlighted_value = totals["grand_total"]
    else:
        final_heading = "FINAL OFFER VALUE" if totals["has_discount"] else "OFFER VALUE"
        highlighted_value = totals["final"]

    set_cell_text(b, final_heading, True, WHITE, 12, WD_ALIGN_PARAGRAPH.CENTER)
    p = b.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(fmt_inr(highlighted_value) + "/-")
    r.bold=True
    r.font.size=Pt(26)
    r.font.name="Aptos Display"
    r.font.color.rgb=RGBColor.from_string(WHITE)
    p = b.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    bottom_note_text = "Freight extra as applicable" if has_gst else "GST & Freight extra as applicable"
    r = p.add_run(bottom_note_text)
    r.font.size=Pt(9.5)
    r.font.name="Aptos"
    r.font.color.rgb=RGBColor.from_string("9CC3E6")

    doc.add_paragraph("")
    note = doc.add_table(rows=1, cols=1)
    c = note.cell(0,0)
    set_cell_shading(c, "F7F9FC")
    set_cell_border(c)
    set_cell_margins(c, 110,120,110,120)
    price_words_amount = totals["grand_total"] if has_gst else totals["final"]
    set_cell_text(c, "Price in Words: " + amount_words(price_words_amount) + ".", True, TEXT, 10.5)
    note_text = clean(static.get("Price Notes"))
    if note_text:
        p = c.add_paragraph(note_text)
        p.runs[0].font.size=Pt(10)
        p.runs[0].font.color.rgb=RGBColor.from_string("6B7280")

    doc.add_page_break()

    # Page 4
    heading(doc, "COMMERCIAL TERMS & CUSTOMER SCOPE", "Concise terms for clear decision-making")
    term_rows = ["Payment Terms", "Delivery", "Warranty", "Price Basis", "GST", "Freight & Insurance", "Commissioning", "Offer Validity"]
    ttable = doc.add_table(rows=1, cols=2)
    set_cell_text(ttable.cell(0,0), "Parameter", True, WHITE, 10.5, WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(ttable.cell(0,1), "Details", True, WHITE, 10.5, WD_ALIGN_PARAGRAPH.CENTER)
    for key in term_rows:
        if key == "GST":
            if has_gst:
                val = f"Included @{totals['gst_pct']:.0%}"
            else:
                val = clean(terms.get(key)) or "Extra @18%"
        else:
            val = clean(terms.get(key))
        if not val:
            continue
        row = ttable.add_row().cells
        set_cell_text(row[0], key, False, TEXT, 10)
        set_cell_text(row[1], val, False, TEXT, 10)
    style_table(ttable)

    doc.add_paragraph("")
    panels = doc.add_table(rows=1, cols=2)
    c1, c2 = panels.row_cells(0)
    for c, fill in [(c1, LIGHT_BLUE), (c2, "F7F9FC")]:
        set_cell_shading(c, fill)
        set_cell_border(c)
        set_cell_margins(c,120,120,120,120)
    set_cell_text(c1, "CUSTOMER SCOPE", True, BLUE, 10.5)
    for i in range(1, 5):
        val = clean(terms.get(f"Customer Scope {i}"))
        if val:
            p = c1.add_paragraph("• " + val)
            p.runs[0].font.size=Pt(10)
            p.runs[0].font.name="Aptos"
    set_cell_text(c2, "EXCEPTIONAL CONDITIONS", True, BLUE, 10.5)
    if clean(terms.get("Exceptional Conditions")):
        p = c2.add_paragraph(clean(terms.get("Exceptional Conditions")))
        p.runs[0].font.size=Pt(10)
        p.runs[0].font.name="Aptos"

    doc.add_page_break()

    # Page 5
    heading(doc, "BANKING, CONTACT & CLOSING", "Order placement and payment details")
    bc = doc.add_table(rows=1, cols=2)
    bc.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, c in enumerate(bc.row_cells(0)):
        set_cell_shading(c, LIGHT_BLUE if idx == 0 else WHITE)
        set_cell_border(c)
        set_cell_margins(c,120,120,120,120)
    c1, c2 = bc.row_cells(0)
    set_cell_text(c1, "BANKING DETAILS", True, BLUE, 10.5)
    for line in [
        settings.get("Company Name"),
        "Bank: " + clean(settings.get("Bank Name")),
        "Account No.: " + clean(settings.get("Account No.")),
        "IFSC: " + clean(settings.get("IFSC")),
        "GSTIN: " + clean(settings.get("GSTIN")),
    ]:
        if clean(line).split(":")[-1].strip():
            p = c1.add_paragraph(clean(line))
            p.runs[0].font.size=Pt(10.2)
            p.runs[0].font.name="Aptos"

    set_cell_text(c2, "ORDER TO BE PLACED ON", True, BLUE, 10.5)
    for line in [
        "M/s " + clean(settings.get("Company Name")),
        settings.get("Address"),
        "Email: " + clean(settings.get("Email")),
        "Mob.: " + clean(settings.get("Mobile")),
    ]:
        if clean(line).split(":")[-1].strip():
            p = c2.add_paragraph(clean(line))
            p.runs[0].font.size=Pt(10.2)
            p.runs[0].font.name="Aptos"

    doc.add_paragraph("")
    thanks = doc.add_table(rows=1, cols=1)
    c = thanks.cell(0,0)
    set_cell_shading(c, NAVY)
    set_cell_border(c, NAVY)
    set_cell_margins(c, 130,120,130,120)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("THANK YOU")
    r.bold=True
    r.font.size=Pt(18)
    r.font.name="Aptos Display"
    r.font.color.rgb=RGBColor.from_string(WHITE)
    if clean(static.get("Thank You Text")):
        p = c.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(clean(static.get("Thank You Text")))
        r.font.size=Pt(10.5)
        r.font.name="Aptos"
        r.font.color.rgb=RGBColor.from_string("D8E2EE")

    doc.add_paragraph("")
    sign = doc.add_table(rows=1, cols=2)
    sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = [("For SMS Controls & Automation", "Authorized Signatory\nName & Seal"), ("CUSTOMER ACCEPTANCE", "Signature / PO Reference")]
    for i, c in enumerate(sign.row_cells(0)):
        set_cell_shading(c, WHITE if i == 0 else "F7F9FC")
        set_cell_border(c)
        set_cell_margins(c, 110,120,110,120)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = c.paragraphs[0]
        r = p.add_run(labels[i][0])
        r.bold=True
        r.font.size=Pt(10.8)
        r.font.name="Aptos"
        r.font.color.rgb=RGBColor.from_string(TEXT if i==0 else "6B7280")
        # Reserved blank area for physical signature / stamp / PO reference.
        # The label remains visible below the blank area.
        p_blank = c.add_paragraph()
        p_blank.paragraph_format.space_before = Pt(20)
        p_blank.paragraph_format.space_after = Pt(10)
        p_blank.add_run(" ")

        p2 = c.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        r = p2.add_run(labels[i][1])
        r.bold=True
        r.font.size=Pt(10.5)
        r.font.name="Aptos"
        r.font.color.rgb=RGBColor.from_string(TEXT if i==0 else "6B7280")
    for row in sign.rows:
        trPr = row._tr.get_or_add_trPr()
        trHeight = OxmlElement("w:trHeight")
        trHeight.set(qn("w:val"), "1550")
        trHeight.set(qn("w:hRule"), "atLeast")
        trPr.append(trHeight)

    for p in doc.paragraphs:
        p.paragraph_format.space_after = Pt(3)
        for run in p.runs:
            if not run.font.name:
                run.font.name = "Aptos"
            if not run.font.size:
                run.font.size = Pt(10.5)

    doc.save(output_path)
    return output_path

if __name__ == "__main__":
    logger = setup_logger("proposal_generator")
    try:
        if len(sys.argv) > 1:
            input_file = Path(sys.argv[1])
        else:
            input_file = Path(__file__).parent / "templates" / "Proposal_Data_Template.xlsx"

        logger.info("Input selected: %s", input_file)
        xlsx = normalize_input_file(input_file, logger=logger)
        logger.info("Normalized input file: %s", xlsx)

        out_dir = Path(__file__).parent / "output"
        out_dir.mkdir(exist_ok=True)
        customer, quotation, *_ = read_excel(xlsx)

        quote_no = clean(quotation.get("Quotation No.", "Quotation")).replace("/", "-").replace("\\", "-")
        cust = re.sub(r"[^A-Za-z0-9]+", "_", clean(customer.get("Customer Name", "Customer")).replace("M/s.", "")).strip("_")[:35]
        doc_type = re.sub(r"[^A-Za-z0-9]+", "_", clean(quotation.get("Document Type", "Commercial Proposal"))).strip("_")
        out = out_dir / f"{quote_no}_{cust}_{doc_type}.docx"

        logger.info("Generating DOCX: %s", out)
        create_docx(xlsx, out)
        logger.info("DOCX generated successfully: %s", out)
        print(out)
    except Exception as e:
        logger.exception("Proposal generation failed")
        raise
