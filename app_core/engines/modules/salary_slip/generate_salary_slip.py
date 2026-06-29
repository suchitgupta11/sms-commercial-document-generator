from pathlib import Path
import sys, re, subprocess, shutil, os
from openpyxl import load_workbook
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image, ImageDraw, ImageFont

BASE_DIR = Path(__file__).parent.resolve()
ASSET_LOGO = BASE_DIR / 'assets' / 'logo_from_doc.png'
PRIMARY='0B2F4F'; PRIMARY2='0E416B'; ACCENT='1C7DB8'
LIGHT='EAF4FB'; PALE='F7FBFE'; WHITE='FFFFFF'; TEXT='172B4D'; MUTED='5B6B82'; BORDER='B7CADC'
GREEN='0F766E'; RED='B42318'; FONT='Arial'

def clean(v): return '' if v is None else str(v).strip()
def parse_float(v):
    if v is None or clean(v)=='': return 0.0
    if isinstance(v,str):
        s=v.replace(',','').replace('Rs.','').replace('INR','').replace('₹','').strip()
        return float(s or 0)
    return float(v)
def money(v): return f'{parse_float(v):,.2f}'
def safe_name(s):
    s=re.sub(r'[^A-Za-z0-9_.-]+','_',clean(s)).strip('_')
    return s or 'Salary_Slip'
def map_sheet(wb,name):
    ws=wb[name]; d={}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and clean(row[0]): d[clean(row[0])] = row[1]
    return d
def table_rows(wb, sheet, cols):
    ws=wb[sheet]; rows=[]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and clean(row[0]): rows.append({cols[i]: row[i] if i < len(row) else None for i in range(len(cols))})
    return rows
ONES=['','One','Two','Three','Four','Five','Six','Seven','Eight','Nine','Ten','Eleven','Twelve','Thirteen','Fourteen','Fifteen','Sixteen','Seventeen','Eighteen','Nineteen']
TENS=['','','Twenty','Thirty','Forty','Fifty','Sixty','Seventy','Eighty','Ninety']
def words_under_1000(n):
    n=int(n); parts=[]
    if n>=100: parts += [ONES[n//100], 'Hundred']; n%=100
    if n>=20: parts.append(TENS[n//10]); n%=10
    if n>0: parts.append(ONES[n])
    return ' '.join(parts)
def amount_words(amount):
    n=int(round(parse_float(amount)))
    if n==0: return 'Indian Rupees Zero Only'
    parts=[]
    for div,label in [(10000000,'Crore'),(100000,'Lakh'),(1000,'Thousand')]:
        q=n//div; n%=div
        if q: parts.append(words_under_1000(q)+' '+label)
    if n: parts.append(words_under_1000(n))
    return 'Indian Rupees ' + ' '.join(parts) + ' Only'
def load_data(xlsx):
    wb=load_workbook(xlsx, data_only=True)
    required=['Company_Details','Employee_Details','Salary_Details','Earnings','Deductions','Settings']
    missing=[s for s in required if s not in wb.sheetnames]
    if missing: raise RuntimeError('Missing required worksheet(s): '+', '.join(missing))
    data={'company':map_sheet(wb,'Company_Details'),'employee':map_sheet(wb,'Employee_Details'),'salary':map_sheet(wb,'Salary_Details'),'settings':map_sheet(wb,'Settings'),'earnings':table_rows(wb,'Earnings',['Component','Rate','Payable']),'deductions':table_rows(wb,'Deductions',['Component','Amount'])}
    show_zero=clean(data['settings'].get('Show Zero Deductions')).lower() in ('yes','true','1','y')
    if not show_zero: data['deductions']=[d for d in data['deductions'] if parse_float(d.get('Amount')) != 0]
    return data

def hex_color(h):
    h=h.replace('#',''); return RGBColor(int(h[:2],16), int(h[2:4],16), int(h[4:6],16))
def set_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill.replace('#',''))
def set_border(cell, color=BORDER, sz='5'):
    tcPr=cell._tc.get_or_add_tcPr(); tcBorders=tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None: tcBorders=OxmlElement('w:tcBorders'); tcPr.append(tcBorders)
    for edge in ('top','left','bottom','right'):
        el=tcBorders.find(qn('w:'+edge))
        if el is None: el=OxmlElement('w:'+edge); tcBorders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),sz); el.set(qn('w:space'),'0'); el.set(qn('w:color'),color.replace('#',''))
def set_margins(cell, top=60, start=80, bottom=60, end=80):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')
def set_table_width(table, width_in):
    tblPr=table._tbl.tblPr
    tblW=tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW=OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(int(width_in*1440))); tblW.set(qn('w:type'),'dxa')
    tblLayout=tblPr.find(qn('w:tblLayout'))
    if tblLayout is None:
        tblLayout=OxmlElement('w:tblLayout'); tblPr.append(tblLayout)
    tblLayout.set(qn('w:type'),'fixed')
    table.autofit=False; table.allow_autofit=False

def set_col_widths(table, widths):
    # Set both visible cell widths and the underlying w:tblGrid.
    # Without tblGrid, LibreOffice may redistribute columns and create huge gaps.
    tbl = table._tbl
    tblGrid = tbl.tblGrid
    if tblGrid is None:
        tblGrid = OxmlElement('w:tblGrid')
        tbl.insert(0, tblGrid)
    for child in list(tblGrid):
        tblGrid.remove(child)
    for w in widths:
        gridCol = OxmlElement('w:gridCol')
        gridCol.set(qn('w:w'), str(int(w*1440)))
        tblGrid.append(gridCol)
    for row in table.rows:
        for i,w in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width=Inches(w)
                tcPr=row.cells[i]._tc.get_or_add_tcPr()
                tcW=tcPr.find(qn('w:tcW'))
                if tcW is None:
                    tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
                tcW.set(qn('w:w'), str(int(w*1440))); tcW.set(qn('w:type'),'dxa')
def row_height(row, inches, rule='atLeast'):
    trPr=row._tr.get_or_add_trPr(); trHeight=trPr.find(qn('w:trHeight'))
    if trHeight is None: trHeight=OxmlElement('w:trHeight'); trPr.append(trHeight)
    trHeight.set(qn('w:val'), str(int(inches*1440))); trHeight.set(qn('w:hRule'), rule)
def para_run(paragraph, text, size=8, bold=False, color=TEXT):
    paragraph.paragraph_format.space_before=Pt(0); paragraph.paragraph_format.space_after=Pt(0); paragraph.paragraph_format.line_spacing=1.0
    r=paragraph.add_run(clean(text) or '-')
    r.font.name=FONT; r._element.rPr.rFonts.set(qn('w:eastAsia'),FONT); r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=hex_color(color)
    return r
def cell_text(cell, text, size=8, bold=False, color=TEXT, align='left'):
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=cell.paragraphs[0]; p.text=''; p.alignment={'left':WD_ALIGN_PARAGRAPH.LEFT,'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT}.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    para_run(p,text,size,bold,color)
def spacer(doc, pts=5):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(pts); p.add_run('')



def _font_candidates(bold=False):
    files = []
    if bold:
        files += [
            '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',
            '/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf',
            '/System/Library/Fonts/Supplemental/Arial Bold.ttf',
        ]
    files += [
        '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
        '/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf',
        '/System/Library/Fonts/Supplemental/Arial.ttf',
    ]
    return files

def _img_font(size, bold=False):
    for f in _font_candidates(bold):
        if Path(f).exists():
            return ImageFont.truetype(f, size)
    return ImageFont.load_default()

def _crop_logo_padding(logo):
    try:
        logo = logo.convert('RGBA')
        alpha = logo.getchannel('A')
        bbox = alpha.getbbox()
        if bbox:
            logo = logo.crop(bbox)
        return logo
    except Exception:
        return logo

def _draw_text_fit(draw, xy, text, max_width, font_size, fill, bold=False, min_size=14):
    text = str(clean(text) or '-')
    size = int(font_size)
    while size >= min_size:
        font = _img_font(size, bold)
        bbox = draw.textbbox((0, 0), text, font=font)
        if bbox[2] - bbox[0] <= max_width:
            draw.text(xy, text, font=font, fill=fill)
            return
        size -= 1
    draw.text(xy, text, font=_img_font(min_size, bold), fill=fill)

def _render_salary_header_image(data, out_path):
    """Render Salary Slip header with the same geometry as Invoice/PI header.

    This prevents Word/LibreOffice from redistributing the header into equal blocks.
    The image is high-resolution and then placed in DOCX at page width.
    """
    company, sal = data['company'], data['salary']
    SCALE = 3
    BW, BH = 1500, 350
    Wpx, Hpx = BW*SCALE, BH*SCALE
    img = Image.new('RGB', (Wpx, Hpx), 'white')
    d = ImageDraw.Draw(img)
    navy = '#0B2F4F'; accent = '#1F78B4'; soft = '#F4F9FD'; text = '#172B4D'; muted = '#5E6C84'; white = '#FFFFFF'
    def sc(v): return int(round(v*SCALE))
    def rect(t): return tuple(sc(v) for v in t)

    d.rounded_rectangle(rect((2,2,BW-3,BH-3)), radius=sc(7), outline=navy, width=sc(2), fill=white)
    d.rectangle(rect((2,2,BW-3,72)), fill=navy)
    company_name = clean(company.get('Company Name')) or 'SMS Controls & Automation'
    top_title = 'SALARY SLIP'
    d.text((sc(32), sc(24)), company_name, font=_img_font(sc(30), True), fill=white)
    tf = _img_font(sc(30), True)
    tb = d.textbbox((0,0), top_title, font=tf)
    d.text((Wpx-sc(32)-(tb[2]-tb[0]), sc(24)), top_title, font=tf, fill=white)

    logo_box = rect((65,105,285,310))
    if ASSET_LOGO.exists():
        try:
            logo = _crop_logo_padding(Image.open(ASSET_LOGO))
            target_w = logo_box[2]-logo_box[0]; target_h = logo_box[3]-logo_box[1]
            scale = min(target_w/logo.width, target_h/logo.height)
            logo = logo.resize((max(1,int(logo.width*scale)), max(1,int(logo.height*scale))), Image.LANCZOS).convert('RGBA')
            lx = logo_box[0] + (target_w-logo.width)//2
            ly = logo_box[1] + (target_h-logo.height)//2
            img.paste(logo, (lx,ly), logo)
        except Exception:
            d.text((sc(105), sc(184)), 'SMS', font=_img_font(sc(62), True), fill=navy)
    else:
        d.text((sc(105), sc(184)), 'SMS', font=_img_font(sc(62), True), fill=navy)

    cx = sc(330)
    _draw_text_fit(d, (cx, sc(104)), company_name, sc(610), sc(43), navy, True, sc(30))
    tagline = clean(company.get('Tagline')) or 'Industrial Automation | AC/DC Drives | Repairs | Retrofitting'
    _draw_text_fit(d, (cx, sc(154)), tagline, sc(610), sc(23), muted, False, sc(18))
    address1 = clean(company.get('Address Line 1')) or 'Office No.2, Plot No.241, Loha Mandi'
    address2 = clean(company.get('Address Line 2')) or 'Ghaziabad, Uttar Pradesh - 201009'
    gst = clean(company.get('GSTIN')) or '09JSMPS2386Q1ZB'
    state = clean(company.get('State')) or clean(company.get('State Name')) or 'Uttar Pradesh'
    state_code = clean(company.get('State Code')) or '09'
    email = clean(company.get('Email')) or 'info@smscontrols.com'
    phone = clean(company.get('Phone')) or clean(company.get('Contact No.')) or clean(company.get('Mobile')) or '+91-8826059159 / 8595231536'
    _draw_text_fit(d, (cx, sc(198)), address1, sc(610), sc(21), text, False, sc(17))
    _draw_text_fit(d, (cx, sc(233)), address2, sc(610), sc(21), text, False, sc(17))
    _draw_text_fit(d, (cx, sc(268)), f'GSTIN: {gst} | State: {state}, Code: {state_code}', sc(610), sc(20), text, False, sc(16))
    _draw_text_fit(d, (cx, sc(303)), f'Email: {email} | Phone: {phone}', sc(610), sc(20), text, False, sc(15))

    # Pay Period card exactly in the Invoice Details-card position.
    card_w, card_h = sc(440), sc(218)
    card_x, card_y = Wpx - sc(72) - card_w, sc(88)
    d.rounded_rectangle((card_x, card_y, card_x+card_w, card_y+card_h), radius=sc(6), outline=accent, width=sc(2), fill=soft)
    d.rectangle((card_x, card_y, card_x+card_w, card_y+sc(58)), fill=navy)
    head = 'PAY PERIOD'
    hf = _img_font(sc(27), True); hb = d.textbbox((0,0), head, font=hf)
    d.text((card_x+(card_w-(hb[2]-hb[0]))//2, card_y+sc(17)), head, font=hf, fill=white)
    rows = [
        ('Month', f"{clean(sal.get('Month'))} {clean(sal.get('Year'))}".strip(), True),
        ('Pay Days', clean(sal.get('Pay Days')) or '0', False),
        ('LOP Days', clean(sal.get('LOP Days')) or '0', False),
        ('Payment', clean(sal.get('Payment Date')) or '-', False),
    ]
    y = card_y + sc(78)
    for lab, val, bold in rows:
        d.text((card_x+sc(38), y), lab, font=_img_font(sc(18), True), fill=muted)
        d.text((card_x+sc(210), y), str(val), font=_img_font(sc(18), bold), fill=navy if bold else text)
        y += sc(32)
    out_path = Path(out_path); out_path.parent.mkdir(parents=True, exist_ok=True)
    img.save(out_path, quality=96)
    return out_path

def _add_invoice_style_salary_header(doc, data):
    tmp_dir = BASE_DIR / '_generated_headers'
    header_path = tmp_dir / f"salary_{safe_name(clean(data['salary'].get('Slip No.')) or 'sample')}_header.png"
    _render_salary_header_image(data, header_path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.add_run().add_picture(str(header_path), width=Inches(7.67))

def init_doc():
    doc=Document(); sec=doc.sections[0]
    # V5.5.6: salary slip is portrait A4. This prevents the UI/Word from scaling a landscape page and scattering text.
    sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.top_margin=Inches(0.28); sec.bottom_margin=Inches(0.25); sec.left_margin=Inches(0.30); sec.right_margin=Inches(0.30)
    for style in doc.styles:
        if style.type == 1:
            style.font.name=FONT; style._element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
    return doc

def create_docx(data, out_path):
    company, emp, sal=data['company'],data['employee'],data['salary']
    earnings=data['earnings']; deductions=data['deductions'] or [{'Component':'No Deductions','Amount':0}]
    gross=sum(parse_float(x.get('Payable')) for x in earnings)
    total_ded=sum(parse_float(x.get('Amount')) for x in data['deductions'])
    net=gross-total_ded
    doc=init_doc(); page_w=7.67

    _add_invoice_style_salary_header(doc, data)
    spacer(doc,4)

    # Employee + Payroll Details: two visually separate premium cards.
    # Implemented as one stable table with a real gutter column, not nested tables,
    # so Word/LibreOffice preserve the layout while the cards look independent.
    left=[('Employee Code',emp.get('Employee Code')),('Employee Name',emp.get('Employee Name')),('Department',emp.get('Department')),('Designation',emp.get('Designation')),('Branch',emp.get('Branch')),('PAN No.',emp.get('PAN No.')),('UAN No.',clean(emp.get('UAN No.')) or 'N/A'),('PF Number',clean(emp.get('PF Number')) or 'N/A')]
    right=[('Bank Name',emp.get('Bank Name')),('Account No.',emp.get('Account No.')),('IFSC Code',emp.get('IFSC Code')),('Slip No.',sal.get('Slip No.')),('Pay Period',sal.get('Pay Period')),('Payment Date',sal.get('Payment Date')),('Pay Days',sal.get('Pay Days')),('LOP Days',sal.get('LOP Days'))]
    detail_rows=max(len(left),len(right))
    details=doc.add_table(rows=detail_rows+1, cols=5)
    details.alignment=WD_TABLE_ALIGNMENT.CENTER
    set_table_width(details,7.35)
    # Left card + controlled gutter + right card; centered as a cohesive block like PDF.
    set_col_widths(details,[1.25,2.30,0.25,1.25,2.30])
    for row in details.rows:
        row_height(row,0.28)
        for cidx,cell in enumerate(row.cells):
            if cidx == 2:
                # Transparent gutter: no borders and white background.
                set_shading(cell,WHITE)
                set_margins(cell,top=0,bottom=0,start=0,end=0)
                tcPr=cell._tc.get_or_add_tcPr()
                tcBorders=tcPr.first_child_found_in('w:tcBorders')
                if tcBorders is None:
                    tcBorders=OxmlElement('w:tcBorders'); tcPr.append(tcBorders)
                for edge in ('top','left','bottom','right'):
                    el=tcBorders.find(qn('w:'+edge))
                    if el is None: el=OxmlElement('w:'+edge); tcBorders.append(el)
                    el.set(qn('w:val'),'nil')
            else:
                set_border(cell,BORDER,'5')
                set_margins(cell,top=46,bottom=46,start=70,end=70)
                set_shading(cell,PALE)
    left_hdr=details.cell(0,0).merge(details.cell(0,1))
    right_hdr=details.cell(0,3).merge(details.cell(0,4))
    for hdr,title in [(left_hdr,'EMPLOYEE DETAILS'),(right_hdr,'PAYROLL DETAILS')]:
        set_shading(hdr,PRIMARY)
        set_border(hdr,PRIMARY,'7')
        cell_text(hdr,title,8.7,True,WHITE,'center')
    for i in range(detail_rows):
        rowidx=i+1
        fill=WHITE if i%2==0 else PALE
        for j in [0,1,3,4]: set_shading(details.cell(rowidx,j),fill)
        if i < len(left):
            lab,val=left[i]
            cell_text(details.cell(rowidx,0),lab,6.9,True,MUTED,'left')
            cell_text(details.cell(rowidx,1),val,6.9, lab == 'Employee Name', TEXT,'left')
        else:
            cell_text(details.cell(rowidx,0),'',6.9,False,TEXT,'left')
            cell_text(details.cell(rowidx,1),'',6.9,False,TEXT,'left')
        if i < len(right):
            lab,val=right[i]
            cell_text(details.cell(rowidx,3),lab,6.9,True,MUTED,'left')
            cell_text(details.cell(rowidx,4),val,6.9, lab == 'Slip No.', TEXT,'left')
        else:
            cell_text(details.cell(rowidx,3),'',6.9,False,TEXT,'left')
            cell_text(details.cell(rowidx,4),'',6.9,False,TEXT,'left')
    spacer(doc,5)

    # Earnings / deductions section with guarded numeric cell widths.
    n=max(len(earnings),len(deductions),4)
    comp=doc.add_table(rows=1+1+n+2, cols=6)
    comp.alignment=WD_TABLE_ALIGNMENT.CENTER
    set_table_width(comp,page_w)
    set_col_widths(comp,[1.55,0.92,0.94,1.65,0.86,1.75])
    for row in comp.rows:
        row_height(row,0.27)
        for cell in row.cells:
            set_border(cell,BORDER,'5'); set_margins(cell,top=42,bottom=42,start=70,end=70)
    eh=comp.cell(0,0).merge(comp.cell(0,2)); dh=comp.cell(0,3).merge(comp.cell(0,5))
    for cell,title in [(eh,'EARNINGS'),(dh,'DEDUCTIONS')]:
        set_shading(cell,PRIMARY); cell_text(cell,title,8.3,True,WHITE,'center')
    for j,hx in enumerate(['Component','Rate','Payable','Deduction','Rate','Amount']):
        set_shading(comp.cell(1,j),PRIMARY2); cell_text(comp.cell(1,j),hx,6.8,True,WHITE,'center')
    for i in range(n):
        rowidx=2+i; fill=WHITE if i%2==0 else PALE
        for j in range(6): set_shading(comp.cell(rowidx,j),fill)
        e=earnings[i] if i < len(earnings) else None; d=deductions[i] if i < len(deductions) else None
        cell_text(comp.cell(rowidx,0),e.get('Component') if e else '',6.8,False,TEXT,'left')
        cell_text(comp.cell(rowidx,1),money(e.get('Rate')) if e else '',6.8,False,TEXT,'right')
        cell_text(comp.cell(rowidx,2),money(e.get('Payable')) if e else '',6.8,False,TEXT,'right')
        cell_text(comp.cell(rowidx,3),d.get('Component') if d else '',6.8,False,TEXT,'left')
        cell_text(comp.cell(rowidx,4),'',6.8,False,TEXT,'right')
        cell_text(comp.cell(rowidx,5),money(d.get('Amount')) if d else '',6.8,False,TEXT,'right')
    totalrow=2+n
    for j in range(6): set_shading(comp.cell(totalrow,j),LIGHT)
    cell_text(comp.cell(totalrow,0),'Gross Earnings',7.2,True,PRIMARY,'left')
    cell_text(comp.cell(totalrow,1),money(gross),7.2,True,PRIMARY,'right')
    cell_text(comp.cell(totalrow,2),money(gross),7.2,True,PRIMARY,'right')
    cell_text(comp.cell(totalrow,3),'Total Deductions',7.2,True,PRIMARY,'left')
    cell_text(comp.cell(totalrow,5),money(total_ded),7.2,True,PRIMARY,'right')
    summary=totalrow+1
    comp.cell(summary,0).merge(comp.cell(summary,1)); comp.cell(summary,2).merge(comp.cell(summary,3)); comp.cell(summary,4).merge(comp.cell(summary,5))
    for idx,(lab,val,fill,col) in enumerate([('Gross Pay','Rs. '+money(gross),LIGHT,PRIMARY),('Total Deduction','Rs. '+money(total_ded),'FDF2F2',RED),('Net Payable','Rs. '+money(net),'E8F8EF',GREEN)]):
        cell=comp.cell(summary,[0,2,4][idx]); set_shading(cell,fill); set_border(cell,ACCENT if idx!=1 else BORDER,'8'); set_margins(cell,top=60,bottom=60,start=70,end=70)
        p=cell.paragraphs[0]; p.text=''; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; para_run(p,lab,7.4,True,MUTED)
        p2=cell.add_paragraph(); p2.alignment=WD_ALIGN_PARAGRAPH.CENTER; para_run(p2,val,10.2,True,col)
    spacer(doc,5)

    aw=doc.add_table(rows=2, cols=2)
    aw.alignment=WD_TABLE_ALIGNMENT.CENTER
    set_table_width(aw,page_w)
    set_col_widths(aw,[1.45,6.22])
    for row in aw.rows:
        row_height(row,0.27)
        for cell in row.cells:
            set_border(cell,BORDER,'5'); set_margins(cell,top=42,bottom=42,start=70,end=70); set_shading(cell,PALE)
    cell_text(aw.cell(0,0),'Amount in Words:',7.1,True,PRIMARY,'left')
    cell_text(aw.cell(0,1),amount_words(net),7.1,True,TEXT,'left')
    note=aw.cell(1,0).merge(aw.cell(1,1))
    cell_text(note,clean(sal.get('Footer Note')) or 'This is a computer generated salary slip and does not require signature.',6.6,False,MUTED,'center')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(0)
    para_run(p,'Computer Generated Salary Slip | SMS Controls & Automation',6.6,False,MUTED)
    doc.save(out_path); return out_path

def _pdf_font_setup():
    """Register dependable TrueType fonts for sharp native PDF output."""
    regular = None
    bold = None
    for f in _font_candidates(False):
        if Path(f).exists():
            regular = f; break
    for f in _font_candidates(True):
        if Path(f).exists():
            bold = f; break
    try:
        if regular:
            pdfmetrics.registerFont(TTFont('SMSFont', regular))
        else:
            return ('Helvetica', 'Helvetica-Bold')
        if bold:
            pdfmetrics.registerFont(TTFont('SMSFont-Bold', bold))
        else:
            pdfmetrics.registerFont(TTFont('SMSFont-Bold', regular))
        return ('SMSFont', 'SMSFont-Bold')
    except Exception:
        return ('Helvetica', 'Helvetica-Bold')

def _pdf_text(c, x, y, txt, size=8, bold=False, color_hex=TEXT, align='left'):
    reg, bld = _pdf_font_setup()
    font = bld if bold else reg
    c.setFont(font, size)
    c.setFillColor(colors.HexColor('#' + color_hex.replace('#','')))
    txt = clean(txt) or '-'
    if align == 'right':
        c.drawRightString(x, y, txt)
    elif align == 'center':
        c.drawCentredString(x, y, txt)
    else:
        c.drawString(x, y, txt)

def _pdf_card(c, x, y_top, w, title, rows, row_h=18, label_w=110):
    """Draw one premium information block matching the DOCX card."""
    header_h = 22
    h = header_h + row_h * len(rows)
    y = y_top - h
    c.setStrokeColor(colors.HexColor('#' + BORDER)); c.setLineWidth(0.8)
    c.setFillColor(colors.HexColor('#' + PALE))
    c.roundRect(x, y, w, h, 4, stroke=1, fill=1)
    c.setFillColor(colors.HexColor('#' + PRIMARY))
    c.roundRect(x, y_top-header_h, w, header_h, 4, stroke=0, fill=1)
    # square off bottom of header so it blends with the card body
    c.rect(x, y_top-header_h, w, header_h/2, stroke=0, fill=1)
    _pdf_text(c, x+w/2, y_top-15, title, 8.2, True, WHITE, 'center')
    yy = y_top - header_h
    for i,(lab,val,bold_val) in enumerate(rows):
        row_y = yy - (i+1)*row_h
        if i % 2 == 0:
            c.setFillColor(colors.white)
        else:
            c.setFillColor(colors.HexColor('#' + PALE))
        c.rect(x, row_y, w, row_h, stroke=0, fill=1)
        c.setStrokeColor(colors.HexColor('#' + BORDER)); c.setLineWidth(0.45)
        c.line(x, row_y, x+w, row_y)
        c.line(x+label_w, row_y, x+label_w, row_y+row_h)
        _pdf_text(c, x+8, row_y+5.2, lab, 6.8, True, MUTED, 'left')
        _pdf_text(c, x+label_w+8, row_y+5.2, val, 6.8, bold_val, TEXT, 'left')
    return y

def _pdf_money_box(c, x, y_top, w, label, value, fill_hex, value_hex):
    h=38; y=y_top-h
    c.setFillColor(colors.HexColor('#'+fill_hex.replace('#','')))
    c.setStrokeColor(colors.HexColor('#'+BORDER)); c.setLineWidth(0.7)
    c.roundRect(x,y,w,h,4,stroke=1,fill=1)
    _pdf_text(c,x+w/2,y+24,label,7.3,True,MUTED,'center')
    _pdf_text(c,x+w/2,y+9,value,10.4,True,value_hex,'center')
    return y

def _pdf_table(c, x, y_top, w, headers, rows, widths, title=None):
    title_h=20 if title else 0
    head_h=18; row_h=16
    h=title_h+head_h+row_h*(len(rows)+1)
    y=y_top-h
    c.setStrokeColor(colors.HexColor('#'+BORDER)); c.setLineWidth(0.55)
    if title:
        c.setFillColor(colors.HexColor('#'+PRIMARY)); c.rect(x, y_top-title_h, w, title_h, stroke=0, fill=1)
        _pdf_text(c,x+w/2,y_top-14,title,7.8,True,WHITE,'center')
    yy=y_top-title_h
    c.setFillColor(colors.HexColor('#'+PRIMARY2)); c.rect(x, yy-head_h, w, head_h, stroke=0, fill=1)
    xx=x
    for hdr,cw in zip(headers,widths):
        _pdf_text(c, xx+cw/2, yy-12, hdr, 6.6, True, WHITE, 'center')
        xx+=cw
    yy-=head_h
    for i,row in enumerate(rows):
        c.setFillColor(colors.white if i%2==0 else colors.HexColor('#'+PALE))
        c.rect(x, yy-row_h, w, row_h, stroke=0, fill=1)
        xx=x
        for j,cw in enumerate(widths):
            c.setStrokeColor(colors.HexColor('#'+BORDER)); c.setLineWidth(0.35)
            c.rect(xx, yy-row_h, cw, row_h, stroke=1, fill=0)
            val = row[j] if j < len(row) else ''
            align='right' if j in (1,2,4,5) else 'left'
            tx = xx+cw-5 if align=='right' else xx+5
            _pdf_text(c, tx, yy-11, val, 6.4, False, TEXT, align)
            xx+=cw
        yy-=row_h
    return y

def create_pdf(data,out_path):
    """Native ReportLab salary slip PDF.

    V6.0.1: no LibreOffice dependency. This ensures the UI creates both DOCX
    and PDF on machines where office conversion is unavailable.
    """
    company, emp, sal=data['company'],data['employee'],data['salary']
    earnings=data['earnings']; deductions=data['deductions'] or [{'Component':'No Deductions','Amount':0}]
    gross=sum(parse_float(x.get('Payable')) for x in earnings)
    total_ded=sum(parse_float(x.get('Amount')) for x in data['deductions'])
    net=gross-total_ded
    out_path=Path(out_path); out_path.parent.mkdir(parents=True,exist_ok=True)
    c=canvas.Canvas(str(out_path), pagesize=A4)
    page_w,page_h=A4
    margin=22
    content_w=page_w-2*margin
    # Header image reused from the DOCX header renderer for exact visual identity.
    hpath=BASE_DIR/'_generated_headers'/f"salary_pdf_{safe_name(clean(sal.get('Slip No.')) or 'sample')}_header.png"
    _render_salary_header_image(data,hpath)
    header_h=content_w*(350/1500)
    c.drawImage(str(hpath), margin, page_h-margin-header_h, width=content_w, height=header_h, preserveAspectRatio=True, mask='auto')
    y=page_h-margin-header_h-14

    left_rows=[('Employee Code',emp.get('Employee Code'),False),('Employee Name',emp.get('Employee Name'),True),('Department',emp.get('Department'),False),('Designation',emp.get('Designation'),False),('Branch',emp.get('Branch'),False),('PAN No.',emp.get('PAN No.'),False),('UAN No.',clean(emp.get('UAN No.')) or 'N/A',False),('PF Number',clean(emp.get('PF Number')) or 'N/A',False)]
    right_rows=[('Bank Name',emp.get('Bank Name'),False),('Account No.',emp.get('Account No.'),False),('IFSC Code',emp.get('IFSC Code'),False),('Slip No.',sal.get('Slip No.'),True),('Pay Period',sal.get('Pay Period'),False),('Payment Date',sal.get('Payment Date'),False),('Pay Days',sal.get('Pay Days'),False),('LOP Days',sal.get('LOP Days'),False)]
    gap=14; card_w=(content_w-gap)/2
    _pdf_card(c, margin, y, card_w, 'EMPLOYEE DETAILS', left_rows, row_h=17.2, label_w=112)
    bottom_cards=_pdf_card(c, margin+card_w+gap, y, card_w, 'PAYROLL DETAILS', right_rows, row_h=17.2, label_w=112)
    y=bottom_cards-12

    box_gap=8; box_w=(content_w-2*box_gap)/3
    _pdf_money_box(c, margin, y, box_w, 'Gross Pay', 'Rs. '+money(gross), LIGHT, PRIMARY)
    _pdf_money_box(c, margin+box_w+box_gap, y, box_w, 'Total Deduction', 'Rs. '+money(total_ded), 'FDF2F2', RED)
    y=_pdf_money_box(c, margin+2*(box_w+box_gap), y, box_w, 'Net Payable', 'Rs. '+money(net), 'E8F8EF', GREEN)-12

    n=max(len(earnings),len(deductions),4)
    e_rows=[]
    for i in range(n):
        e=earnings[i] if i<len(earnings) else None; d=deductions[i] if i<len(deductions) else None
        e_rows.append([
            e.get('Component') if e else '', money(e.get('Rate')) if e else '', money(e.get('Payable')) if e else '',
            d.get('Component') if d else '', '', money(d.get('Amount')) if d else ''
        ])
    e_rows.append(['Gross Earnings', money(gross), money(gross), 'Total Deductions', '', money(total_ded)])
    table_w=content_w
    widths=[table_w*0.23, table_w*0.12, table_w*0.13, table_w*0.25, table_w*0.10, table_w*0.17]
    y=_pdf_table(c, margin, y, table_w, ['Earnings Component','Rate','Payable','Deduction Component','Rate','Amount'], e_rows, widths, title=None)-10

    # Amount in words block
    aw_h=28
    c.setFillColor(colors.HexColor('#'+PALE)); c.setStrokeColor(colors.HexColor('#'+BORDER)); c.setLineWidth(0.5)
    c.rect(margin, y-aw_h, content_w, aw_h, stroke=1, fill=1)
    c.line(margin+112, y-aw_h, margin+112, y)
    _pdf_text(c, margin+8, y-17, 'Amount in Words:', 6.8, True, PRIMARY, 'left')
    _pdf_text(c, margin+120, y-17, amount_words(net), 6.8, True, TEXT, 'left')
    y-=aw_h+10
    _pdf_text(c, page_w/2, y, clean(sal.get('Footer Note')) or 'This is a computer generated salary slip and does not require signature.', 6.4, False, MUTED, 'center')
    _pdf_text(c, page_w/2, y-12, 'Computer Generated Salary Slip | SMS Controls & Automation', 6.2, False, MUTED, 'center')
    c.showPage(); c.save()
    return out_path

def generate(xlsx, output_dir):
    data=load_data(xlsx); out_dir=Path(output_dir); out_dir.mkdir(parents=True,exist_ok=True)
    sal=data['salary']; emp=data['employee']; base=f"{safe_name(clean(sal.get('Slip No.')) or 'SAL')}_{safe_name(clean(emp.get('Employee Name')))}_Salary_Slip"
    docx=out_dir/(base+'.docx'); pdf=out_dir/(base+'.pdf')
    create_docx(data,docx)
    pdf_created=False
    soffice=shutil.which('soffice') or shutil.which('libreoffice')
    if soffice:
        try:
            if pdf.exists(): pdf.unlink()
            env=os.environ.copy(); env.setdefault('HOME', str(out_dir.resolve()))
            subprocess.run([soffice,'--headless','--convert-to','pdf','--outdir',str(out_dir),str(docx)],capture_output=True,text=True,env=env)
            pdf_created=pdf.exists()
        except Exception: pdf_created=False
    if not pdf_created: create_pdf(data,pdf)
    return [docx,pdf]
if __name__=='__main__':
    if len(sys.argv)<2:
        print('Usage: python generate_salary_slip.py <template.xlsx> [output_dir]', file=sys.stderr); sys.exit(1)
    files=generate(Path(sys.argv[1]), Path(sys.argv[2]) if len(sys.argv)>2 else BASE_DIR/'output')
    for f in files: print(f)
