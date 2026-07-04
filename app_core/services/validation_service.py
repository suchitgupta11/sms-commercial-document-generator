from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook


def validate_required_sheets(xlsx_path: str | Path, required_sheets: list[str]) -> list[str]:
    wb = load_workbook(Path(xlsx_path), read_only=True, data_only=True)
    available = set(wb.sheetnames)
    return [sheet for sheet in required_sheets if sheet not in available]
