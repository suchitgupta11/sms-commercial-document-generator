from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook

from app_core.services.registry_service import DOCUMENTS, DocumentDefinition


class TemplateService:
    def infer_document_type(self, xlsx_path: Path) -> str | None:
        try:
            wb = load_workbook(xlsx_path, read_only=True, data_only=True)
            available = set(wb.sheetnames)
        except Exception:
            return None
        filename = str(xlsx_path).lower()
        if {"Company_Details", "Employee_Details", "Salary_Details", "Earnings", "Deductions"}.issubset(available):
            return "Salary Slip"
        if {"Customer_Details", "Quotation_Details", "Items"}.issubset(available):
            return "Quotation"
        if {"Company_Details", "Customer_Details", "Challan_Details", "Items"}.issubset(available):
            return "Delivery Challan"
        if {"Company_Details", "Bank_Details", "Customer_Details", "Invoice_Details", "Items"}.issubset(available):
            if any(t in filename for t in ("purchase_order", "purchase-order", "_po", "po_")):
                return "Purchase Order"
            return "Proforma Invoice" if any(t in filename for t in ("proforma", "pi_", "_pi")) else "Tax Invoice"
        return None

    def validate(self, document: DocumentDefinition, xlsx_path: Path) -> DocumentDefinition:
        try:
            wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        except Exception as exc:
            raise RuntimeError(f"Unable to read selected spreadsheet. Please check the file.\n\nReason: {exc}") from exc
        available = set(wb.sheetnames)
        missing = [sheet for sheet in document.required_sheets if sheet not in available]
        if missing:
            detected = self.infer_document_type(xlsx_path)
            if detected and detected != document.display_name:
                return DOCUMENTS[detected]
            raise RuntimeError(
                f"The selected template does not match {document.display_name}.\n\n"
                f"Missing worksheet(s): {', '.join(missing)}\n\n"
                f"Available worksheet(s): {', '.join(wb.sheetnames)}"
            )
        return document
